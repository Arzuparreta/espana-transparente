import os
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock
import openpyxl
import pytest
from presupuestos.autonomico import TERRITORIES, CODES, parse_workbook, replace_records


def workbook(change=None):
    w=openpyxl.Workbook();s=w.active;s.title='DescargaEconomicaDC'
    for _ in range(48):s.append([None]*6)
    s['A8']='Ingresos y Gastos por capítulos. Datos consolidados'
    s['A9']='Liquidación de Presupuestos de las Comunidades Autónomas. Ejercicio 2024'
    s['A10']=TERRITORIES['01'][2];s['E12']='Miles de Euros'
    for i,v in enumerate(['Capítulos de gastos','Presupuesto Inicial','Presupuesto Definitivo','Obligaciones Reconocidas Netas','Pagos Líquidos'],1):s.cell(33,i,v)
    facts={c:[10,12,8,6] for c in '123456789'};facts['5']=[10,12,None,None]
    for code,parts in [('NF','1234567'),('FIN','89'),('TOTAL',['NF','FIN'])]:facts[code]=[sum(facts[c][i] or 0 for c in parts) for i in range(4)]
    for row,code in zip([35,36,37,38,39,41,42,45,46,44,47,48],CODES):
        s.cell(row,1,f'{code}. Capítulo' if code.isdigit() else {'NF':'Operaciones no financieras','FIN':'Operaciones financieras','TOTAL':'Total Gastos'}[code])
        for i,v in enumerate(facts[code],2):s.cell(row,i,v)
    if change:change(s)
    b=BytesIO();w.save(b);return b.getvalue()


def test_definitive_consolidated_stages_and_source_identity():
    r=parse_workbook(workbook(),'01')
    assert len(r)==12
    assert next(v for v in r if v[2]=='5')[5:7]==(None,None)
    assert 'cdcdad=01&ano=2024' in r[0][8]
    assert r[0][:3]==('ANDALUCIA',2024,'1')


@pytest.mark.parametrize('cell,value',[('A8','Sin consolidar'),('A9','Ejercicio 2025'),('A10','Otra comunidad'),('E12','Euros'),('E33','Compromisos'),('A35','2. Capítulo'),('B35',-1),('D35','=1+1'),('B48',999)])
def test_layout_and_meaning_changes_fail_closed(cell,value):
    with pytest.raises(ValueError):parse_workbook(workbook(lambda s:s.cell(s[cell].row,s[cell].column,value)),'01')


def test_partial_replacement_does_not_write():
    cur=Mock()
    with pytest.raises(ValueError):replace_records(cur,parse_workbook(workbook(),'01'))
    cur.execute.assert_not_called()


@pytest.mark.skipif(not os.getenv('TEST_COFOG_DB_URL'),reason='isolated PostgreSQL only')
def test_seed_idempotence_atomic_rollback_and_permissions():
    import psycopg2
    conn=psycopg2.connect(os.environ['TEST_COFOG_DB_URL']);conn.set_client_encoding('UTF8')
    root=Path(__file__).resolve().parents[2]/'supabase/migrations'
    try:
        with conn.cursor() as cur:
            for role in ('anon','authenticated','service_role'):
                cur.execute('SELECT 1 FROM pg_roles WHERE rolname=%s',(role,))
                if not cur.fetchone():cur.execute(f'CREATE ROLE {role}')
            base=(root/'20260514180000_multilevel_responsibility.sql').read_text().split('CREATE TABLE IF NOT EXISTS etl_runs (',1)[1].split(');',1)[0]
            cur.execute('CREATE TABLE IF NOT EXISTS etl_runs ('+base+');')
            cur.execute('DROP TABLE IF EXISTS autonomic_liquidation')
            cur.execute((root/'20261012020000_autonomic_liquidation.sql').read_text())
            cur.execute("SELECT credit,obligations,payments FROM autonomic_liquidation WHERE territory='ANDALUCIA' AND chapter='NF'")
            from decimal import Decimal
            assert cur.fetchone()==tuple(map(Decimal,['47363032.9','42417589.99','41424022.37']))
            cur.execute('SELECT territory,year,chapter,initial,credit,obligations,payments,unit,raw_data FROM autonomic_liquidation')
            import json
            records=[(*r[:8],json.dumps(r[8])) for r in cur.fetchall()]
            replace_records(cur,records);replace_records(cur,records)
            cur.execute('SELECT count(*) FROM autonomic_liquidation');assert cur.fetchone()[0]==204
        conn.commit()
        bad=list(records[100]);bad[3]=-1
        with pytest.raises(psycopg2.Error):
            with conn.cursor() as cur:replace_records(cur,[*records[:100],bad,*records[101:]])
        conn.rollback()
        with conn.cursor() as cur:
            cur.execute('SET LOCAL ROLE anon');cur.execute('SELECT count(*) FROM autonomic_liquidation');assert cur.fetchone()[0]==204
        with pytest.raises(psycopg2.Error):
            with conn.cursor() as cur:cur.execute('DELETE FROM autonomic_liquidation')
        conn.rollback()
    finally:conn.close()


def test_download_failure_never_returns_a_partial_comparison(monkeypatch):
    from presupuestos import autonomico
    client=Mock();client.__enter__=Mock(return_value=client);client.__exit__=Mock(return_value=False)
    client.get.return_value=Mock(text='<input name="__VIEWSTATE" value="v"><input name="__EVENTVALIDATION" value="e">')
    first=Mock(content=workbook(),url=autonomico.source_url('01'))
    client.post.side_effect=[first,RuntimeError('source unavailable')]
    monkeypatch.setattr(autonomico.httpx,'Client',lambda **kwargs:client)
    with pytest.raises(RuntimeError,match='unavailable'):autonomico.fetch_records()
