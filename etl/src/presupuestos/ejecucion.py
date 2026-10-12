"""IGAE state expenditure budget: current-budget execution, bounded 2025 monthly series."""
import argparse
import calendar
import hashlib
import json
import math
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

import httpx
import openpyxl

SCOPE = "ES_STATE_BUDGET"
PERIOD = "2025-12"
SOURCE_URL = "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20DICIEMBRE%202025.xlsx"
PAGE_URL = "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Paginas/imejecucionpresupuesto2025.aspx"
SOURCES = {
    "2025-01": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20ENERO%202025%20(EXCEL).xlsx",
    "2025-02": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20FEBRERO%202025%20(EXCEL).xlsx",
    "2025-03": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20MARZO%202025%20(EXCEL).xlsx",
    "2025-04": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20ABRIL%202025%20(EXCEL).xlsx",
    "2025-05": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20MAYO%202025%20(EXCEL).xlsx",
    "2025-06": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20JUNIO%202025%20(EXCEL).xlsx",
    "2025-07": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20JULIO%202025%20(EXCEL).xlsx",
    "2025-08": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20AGOSTO%202025%20(EXCEL).xlsx",
    "2025-09": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20SEPTIEMBRE%202025%20(EXCEL).xlsx",
    "2025-10": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20OCTUBRE%202025%20(EXCEL).xlsx",
    "2025-11": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20NOVIEMBRE%202025.xlsx",
    "2025-12": "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadPublica/CPE/EjecucionPresupuestaria/Documents/MENSUAL%20DICIEMBRE%202025.xlsx"
}
FIELDS = ["credit", "committed", "obligations", "payments"]
CODES = [str(i) for i in range(1, 10)] + ["NF", "FIN", "TOTAL"]
MAPPING = "igae-current-chapters-v1"
MONTHS = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def norm(value):
    return re.sub(r"[^A-Z0-9]", "", unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper())


def amount(value):
    if value == "-": return None  # preserve source notation, never manufacture zero
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError("Invalid IGAE amount or uncached formula")
    return Decimal(str(value))


def code(label):
    label = norm(label)
    if label in ("TOTALES", "TOTAL"): return "TOTAL"
    if label == "TOTALOPERACIONESNOFINANCIERAS": return "NF"
    if label == "TOTALOPERACIONESFINANCIERAS": return "FIN"
    match = re.match(r"^([1-9])[^0-9]", label)
    return match[1] if match else None


def validate(rows):
    if set(rows) != set(CODES): raise ValueError("Missing IGAE chapters or totals")
    for group, parts in [("NF", list("1234567")), ("FIN", list("89")), ("TOTAL", ["NF", "FIN"])]:
        for i in range(4):
            target = rows[group][i]
            if target is None: raise ValueError("Missing published execution total")
            known = sum((rows[c][i] for c in parts if rows[c][i] is not None), Decimal(0))
            if abs(known-target) > Decimal("1"):
                raise ValueError(f"IGAE reconciliation failed {group}/{FIELDS[i]}")


def parse_workbook(payload, period=PERIOD, source_url=SOURCE_URL, checked_at=None):
    if not re.fullmatch(r"20[0-9]{2}-(0[1-9]|1[0-2])", period) or len(payload) > 20_000_000:
        raise ValueError("Invalid IGAE period or workbook size")
    year, month = map(int, period.split("-"))
    cut = f"HASTAEL{calendar.monthrange(year,month)[1]}DE{MONTHS[month-1]}"
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        result = {}
        for name, title, headers, width in [
            ("GTOS 004", "ESTADODEEJECUCIONCAPITULOS", ["CREDITOSDEFINITIVOS1", "GASTOSCOMPROMETIDOS2", "OBLIGACIONESRECONOCIDASNETAS3", "PAGOSREALIZADOS4"], 4),
            ("GTOS 013", "PAGOSAPLICADOSAPRESUPUESTOSCAPITULOS", ["PRESUPUESTOCORRIENTE1", "PRESUPUESTOSCERRADOS2", "TOTAL3"], 3),
        ]:
            sheet = workbook[name]
            if sheet.max_row > 100 or sheet.max_column > 30: raise ValueError("Unexpected IGAE layout")
            values = list(sheet.values)
            if "MILESDEEUROS" not in [norm(v) for v in values[0]] or title not in norm(values[1][0]) or cut not in norm(values[1][0]) or str(values[1][1]) != str(year):
                raise ValueError("IGAE title, cut, year or unit changed")
            if [norm(v) for v in values[2][1:width+1]] != headers:
                raise ValueError("IGAE execution columns changed")
            rows = {}
            for row in values[3:]:
                c = code(row[0])
                if c:
                    if c in rows: raise ValueError("Duplicate IGAE chapter")
                    rows[c] = [amount(v) for v in row[1:width+1]]
            result[name] = rows
        rows = result["GTOS 004"]
        validate(rows)
        payments = result["GTOS 013"]
        if set(payments) != set(CODES): raise ValueError("Missing payment reconciliation rows")
        for c in CODES:
            if rows[c][3] != payments[c][0]: raise ValueError("Current-budget payment mismatch")
            if payments[c][2] is not None and abs(sum((v for v in payments[c][:2] if v is not None), Decimal(0))-payments[c][2]) > Decimal("1"):
                raise ValueError("Closed/current payment reconciliation failed")
        metadata = dict(source="IGAE", source_url=source_url, page_url=PAGE_URL, sheet="GTOS 004", unit="MILES_EUR",
            payment_basis="current_budget", mapping_version=MAPPING, retrieved_at=checked_at or datetime.now(timezone.utc).isoformat(),
            workbook_sha256=hashlib.sha256(payload).hexdigest(), period=period)
        return [(SCOPE, period, c, *rows[c], "MILES_EUR", json.dumps(metadata)) for c in CODES]
    finally:
        workbook.close()


def replace_snapshot(cur, records):
    period = records[0][1] if records else None
    if period not in SOURCES or len(records) != 12 or {r[2] for r in records} != set(CODES) or any(r[0] != SCOPE or r[1] != period for r in records):
        raise ValueError("Refusing incomplete or out-of-scope execution replacement")
    cur.execute("DELETE FROM budget_execution WHERE scope=%s AND period=%s", (SCOPE, period))
    cur.executemany("""INSERT INTO budget_execution
        (scope,period,chapter,credit,committed,obligations,payments,unit,raw_data)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)""", records)


def fetch_records():
    # Validate the entire manifest before any database mutation. Each report is
    # an independent cumulative snapshot, never a flow to add to other months.
    records = []
    with httpx.Client(timeout=45, follow_redirects=True) as client:
        for period, url in SOURCES.items():
            response = client.get(url)
            response.raise_for_status()
            records.extend(parse_workbook(response.content, period, url))
    return records


def replace_series(cur, records):
    grouped = {period: [r for r in records if r[1] == period] for period in SOURCES}
    if len(records) != len(SOURCES) * 12 or any(len(rows) != 12 or {r[2] for r in rows} != set(CODES)
        or any(r[0] != SCOPE for r in rows) for rows in grouped.values()):
        raise ValueError("Refusing incomplete execution series")
    for rows in grouped.values():
        replace_snapshot(cur, rows)


def run(dry_run=False):
    if dry_run:
        print(f"Validated {len(fetch_records())} IGAE execution rows; no database writes")
        return
    from common.db import get_pg_conn
    from common.etl_runs import start_run, finish_run
    conn, run_id = get_pg_conn(), None
    try:
        with conn.cursor() as cur: run_id = start_run(cur, pipeline="presupuestos.ejecucion")
        conn.commit()
        records = fetch_records()
        with conn.cursor() as cur:
            replace_series(cur, records)
            finish_run(cur, run_id=run_id, status="succeeded", rows_read=len(records), rows_inserted=len(records))
        conn.commit()
    except Exception as exc:
        conn.rollback()
        if run_id:
            with conn.cursor() as cur: finish_run(cur, run_id=run_id, status="failed", error_summary=str(exc)[:500])
            conn.commit()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    run(parser.parse_args().dry_run)
